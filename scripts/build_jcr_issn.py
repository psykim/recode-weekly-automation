#!/usr/bin/env python3
"""JCR 엑셀(data/JCR2025_202506.xlsx)에서 ISSN·eISSN → JIF 표(data/jcr2025_issn.json)를 만든다.

보고서 생성은 이 JSON 만 읽는다(Actions 에 openpyxl 을 깔지 않으려고). JCR 을 새로 받으면
엑셀을 바꾸고 이 스크립트를 한 번 돌려 JSON 을 커밋한다:  pip install openpyxl && python build_jcr_issn.py
"""
import json
import os
import sys

import openpyxl

HERE = os.path.dirname(__file__)
SRC = os.path.join(HERE, '..', 'data', sys.argv[1] if len(sys.argv) > 1 else 'JCR2025_202506.xlsx')
OUT = os.path.join(HERE, '..', 'data', 'jcr2025_issn.json')

ws = openpyxl.load_workbook(SRC, read_only=True).active
header = [str(c).strip() for c in next(ws.iter_rows(max_row=1, values_only=True))]
col = {name: header.index(name) for name in ('Journal name', 'ISSN', 'eISSN')}
jif_col = next(i for i, h in enumerate(header) if h.endswith('JIF'))

issn = {}
for row in ws.iter_rows(min_row=2, values_only=True):
    try:
        jif = float(row[jif_col])
    except (TypeError, ValueError):
        continue
    for key in (row[col['ISSN']], row[col['eISSN']]):
        if key and str(key).strip().upper() not in ('N/A', ''):
            k = str(key).strip().upper()
            issn[k] = max(jif, issn.get(k, 0.0))   # 같은 ISSN 이 여러 분야에 나오면 값은 같다

with open(OUT, 'w', encoding='utf-8') as f:
    json.dump({'source': os.path.basename(SRC), 'jif_column': header[jif_col], 'issn': issn}, f, ensure_ascii=False, sort_keys=True)
print(f'{len(issn)}개 ISSN → {OUT}')
